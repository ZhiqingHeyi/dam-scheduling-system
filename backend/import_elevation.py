import openpyxl
import pymysql

conn = pymysql.connect(
    host='192.168.1.88',
    port=3306,
    user='root',
    password='!Tmhc20170717',
    database='qbt',
    charset='utf8mb4'
)

try:
    with conn.cursor() as cursor:
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS dam_elevation (
                id INT AUTO_INCREMENT PRIMARY KEY,
                dam_section INT NOT NULL COMMENT '坝段号',
                layer_no INT NOT NULL COMMENT '层号',
                top_elevation DECIMAL(10,2) NOT NULL COMMENT '仓顶高程(m)',
                bottom_elevation DECIMAL(10,2) GENERATED ALWAYS AS (top_elevation - 3) STORED COMMENT '仓底高程(m)',
                UNIQUE KEY uk_dam_layer (dam_section, layer_no)
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='大坝各坝段各层高程数据'
        """)
        print("Table dam_elevation created (or already exists)")

        cursor.execute("SELECT COUNT(*) FROM dam_elevation")
        count = cursor.fetchone()[0]
        if count > 0:
            print(f"Table already has {count} records, skipping import")
        else:
            wb = openpyxl.load_workbook(
                '项目部预期的排仓结果展示形式/2.大坝仓位总排仓计划表(2025.11.28陈总发)(1).xlsx',
                data_only=True
            )
            ws = wb.active

            dam_cols = {}
            for col_idx in range(1, ws.max_column + 1):
                header = ws.cell(row=1, column=col_idx).value
                if header is not None and '坝段' in str(header):
                    dam_id = int(str(header).replace('坝段', ''))
                    dam_cols[dam_id] = col_idx

            records = []
            for dam_id in sorted(dam_cols.keys()):
                col = dam_cols[dam_id]
                layers = []
                for row_idx in range(2, ws.max_row + 1):
                    elev_cell = ws.cell(row=row_idx, column=col)
                    if elev_cell.value is not None:
                        try:
                            layers.append(float(elev_cell.value))
                        except:
                            pass
                layers.reverse()
                for i, elev in enumerate(layers):
                    layer_id = i + 1
                    records.append((dam_id, layer_id, round(elev, 2)))

            cursor.executemany(
                "INSERT INTO dam_elevation (dam_section, layer_no, top_elevation) VALUES (%s, %s, %s)",
                records
            )
            conn.commit()
            print(f"Imported {len(records)} elevation records into dam_elevation table")

    with conn.cursor() as cursor:
        cursor.execute("SELECT COUNT(*) FROM dam_elevation")
        print(f"Total records in dam_elevation: {cursor.fetchone()[0]}")

        for dam_id in [16, 18, 21]:
            cursor.execute(
                "SELECT layer_no, top_elevation, bottom_elevation FROM dam_elevation WHERE dam_section=%s ORDER BY layer_no LIMIT 3",
                (dam_id,)
            )
            rows = cursor.fetchall()
            print(f"\nDam {dam_id} first 3 layers:")
            for row in rows:
                print(f"  L{row[0]}: top={row[1]}, bottom={row[2]}")

finally:
    conn.close()
